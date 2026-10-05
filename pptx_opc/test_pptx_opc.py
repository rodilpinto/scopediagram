# -*- coding: utf-8 -*-
"""Testes do pptx_opc. Viajam com a pasta: nao importam nada de app, nao fazem rede.
As apresentacoes sao geradas em memoria com python-pptx; a "parte de diagrama" (SmartArt) e
injetada a mao, como o PowerPoint a guarda (ppt/diagrams/dataN.xml ligado ao slide).

Rodar da pasta que CONTEM pptx_opc/:
    python -m pytest pptx_opc -q
O teste do render com o PowerPoint de verdade so roda com NUATI_TESTE_POWERPOINT=1 (Windows + PowerPoint).
"""
from __future__ import annotations

import io
import os
import zipfile

import pytest

import pptx_opc
from pptx_opc.opc import CT_NS, Package, _qn

DIAGRAMA_CT = "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"
DIAGRAMA_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"


def _apresentacao(*titulos: str) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for t in titulos:
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _titulos(dados: bytes) -> list[str]:
    from pptx import Presentation
    return [s.shapes.title.text for s in Presentation(io.BytesIO(dados)).slides]


def _com_diagrama(pkg: Package, slide: str, conteudo: bytes = b"<dgm>A</dgm>") -> str:
    """Liga ao slide uma parte de diagrama propria (como um SmartArt) e uma nota."""
    parte = "ppt/diagrams/data1.xml"
    pkg.set_part(parte, conteudo)
    pkg.add_override(parte, DIAGRAMA_CT)
    pkg.set_part("ppt/notesSlides/notesSlide1.xml", b"<notas/>")
    rels = pkg.rels(slide)
    rels.append({"Id": "rId90", "Type": DIAGRAMA_REL, "Target": "../diagrams/data1.xml", "TargetMode": None})
    rels.append({"Id": "rId91", "Type": "notes", "Target": "../notesSlides/notesSlide1.xml", "TargetMode": None})
    pkg.write_rels(slide, rels)
    return parte


def test_versao():
    assert pptx_opc.__version__.count(".") == 2


def test_ida_e_volta_sem_mudanca_abre_no_python_pptx():
    pkg = Package(_apresentacao("A", "B"))
    dados = pkg.to_bytes()
    assert _titulos(dados) == ["A", "B"]
    assert zipfile.ZipFile(io.BytesIO(dados)).namelist()[0] == "[Content_Types].xml"


def test_slides_na_ordem_de_apresentacao():
    pkg = Package(_apresentacao("A", "B", "C"))
    assert pkg.slide_names() == ["ppt/slides/slide1.xml", "ppt/slides/slide2.xml", "ppt/slides/slide3.xml"]


def test_apagar_slide():
    pkg = Package(_apresentacao("A", "B", "C"))
    pkg.delete_slide("ppt/slides/slide2.xml")
    assert not pkg.has_part("ppt/slides/slide2.xml")
    assert pkg.content_type_of("ppt/slides/slide2.xml") is None
    assert _titulos(pkg.to_bytes()) == ["A", "C"]


def test_reordenar():
    pkg = Package(_apresentacao("A", "B", "C"))
    s = pkg.slide_names()
    pkg.reorder([s[2], s[0], s[1]])
    assert _titulos(pkg.to_bytes()) == ["C", "A", "B"]


def test_clonar_slide_vai_para_o_fim_e_abre():
    pkg = Package(_apresentacao("A", "B"))
    novo = pkg.clone_slide("ppt/slides/slide1.xml")
    assert novo == "ppt/slides/slide3.xml"
    assert _titulos(pkg.to_bytes()) == ["A", "B", "A"]


def test_clone_ganha_copia_propria_do_diagrama_e_perde_a_nota():
    pkg = Package(_apresentacao("A"))
    original = _com_diagrama(pkg, "ppt/slides/slide1.xml")
    novo = pkg.clone_slide("ppt/slides/slide1.xml")
    alvos = {r["Target"] for r in pkg.rels(novo)}
    assert "../diagrams/data2.xml" in alvos                      # parte nova, nao a do original
    assert not any("notesSlides" in a for a in alvos)             # nota nao vai para o clone
    assert pkg.part("ppt/diagrams/data2.xml") == pkg.part(original)
    assert pkg.content_type_of("ppt/diagrams/data2.xml") == DIAGRAMA_CT
    pkg.set_part("ppt/diagrams/data2.xml", b"<dgm>B</dgm>")      # mexer no clone nao mexe no original
    assert pkg.part(original) == b"<dgm>A</dgm>"
    assert _titulos(pkg.to_bytes()) == ["A", "A"]


def test_apagar_slide_leva_junto_as_partes_proprias():
    pkg = Package(_apresentacao("A", "B"))
    parte = _com_diagrama(pkg, "ppt/slides/slide1.xml")
    pkg.delete_slide("ppt/slides/slide1.xml")
    assert not pkg.has_part(parte) and not pkg.has_part("ppt/notesSlides/notesSlide1.xml")
    assert pkg.content_type_of(parte) is None
    ct = pkg._xml("[Content_Types].xml")
    assert all(o.get("PartName") != "/" + parte for o in ct.findall(_qn(CT_NS, "Override")))
    assert _titulos(pkg.to_bytes()) == ["B"]


def test_relpath_e_resolve():
    assert Package.resolve("ppt/slides/slide1.xml", "../diagrams/data1.xml") == "ppt/diagrams/data1.xml"
    assert Package.relpath("ppt/slides/slide1.xml", "ppt/diagrams/data1.xml") == "../diagrams/data1.xml"
    assert Package._next_rid([{"Id": "rId1"}, {"Id": "rId3"}]) == "rId2"


@pytest.mark.skipif(os.environ.get("NUATI_TESTE_POWERPOINT") != "1",
                    reason="abre o PowerPoint de verdade: rode com NUATI_TESTE_POWERPOINT=1")
def test_render_com_powerpoint_de_verdade(tmp_path):
    from pptx_opc.render_powerpoint import render
    arquivo = tmp_path / "t.pptx"
    arquivo.write_bytes(_apresentacao("Um", "Dois"))
    pngs = render(str(arquivo), str(tmp_path / "png"))
    assert [os.path.basename(p) for p in pngs] == ["slide01.png", "slide02.png"]
    assert all(os.path.getsize(p) > 1000 for p in pngs)
